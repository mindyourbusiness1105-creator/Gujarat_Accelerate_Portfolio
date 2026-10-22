import hmac, os, re, sqlite3, threading, time
from datetime import datetime
from functools import wraps
from urllib.parse import quote

import requests
from dotenv import load_dotenv
from flask import Flask, abort, g, jsonify, redirect, render_template, request, session
from openpyxl import load_workbook

load_dotenv()
app = Flask(__name__)
SECRET_KEY = os.getenv("SECRET_KEY")
ADMIN_PASSWORD = os.getenv("ADMIN_PASSWORD")
IS_RENDER = os.getenv("RENDER") == "true"
if IS_RENDER and (not SECRET_KEY or not ADMIN_PASSWORD):
    raise RuntimeError("Set SECRET_KEY and ADMIN_PASSWORD before starting in production.")
app.secret_key = SECRET_KEY or "dev-only-change-me"
app.config.update(
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE="Lax",
    SESSION_COOKIE_SECURE=IS_RENDER,
    MAX_CONTENT_LENGTH=5 * 1024 * 1024,
)

ADMIN_PASSWORD = ADMIN_PASSWORD or "admin123"
SITE_URL = (os.getenv("SITE_URL") or os.getenv("RENDER_EXTERNAL_URL") or "http://localhost:5000").rstrip("/")
COMMUNITY_LINK = os.getenv("COMMUNITY_LINK", "")
WA_TOKEN, WA_PHONE_ID = os.getenv("WA_TOKEN"), os.getenv("WA_PHONE_ID")
WA_TEMPLATE, WA_LANG = os.getenv("WA_TEMPLATE", "network_update"), os.getenv("WA_LANG", "en")
DB = os.getenv("DATABASE_PATH", os.path.join(os.path.dirname(os.path.abspath(__file__)), "data.db"))
REGIONS = ["North", "South", "East", "West", "Central"]
KINDS = {"sell": "We sell", "buy": "Looking to buy", "job": "Hiring", "event": "Job fair"}
TAB = {"sell": "buysell", "buy": "buysell", "job": "jobs", "event": "jobs"}

SCHEMA = """
create table if not exists companies(id integer primary key, name text unique collate nocase, sector text, city text,
  region text, state text, about text, phone text, website text, optin integer default 0, active integer default 1);
create table if not exists posts(id integer primary key, company_id integer, kind text, title text, details text,
  extra text, created text, active integer default 1);
create table if not exists notif(id integer primary key, post_id integer, phone text, status text, detail text, at text);
"""


def init_db():
    c = sqlite3.connect(DB)
    c.executescript(SCHEMA)
    c.commit()
    c.close()


def db():
    if "db" not in g:
        g.db = sqlite3.connect(DB)
        g.db.row_factory = sqlite3.Row
    return g.db


@app.teardown_appcontext
def close_db(_):
    d = g.pop("db", None)
    if d:
        d.close()


def now():
    return datetime.now().strftime("%Y-%m-%d %H:%M")


def norm(p):
    d = re.sub(r"\D", "", str(p or ""))
    if len(d) == 10:
        d = "91" + d
    return d if 11 <= len(d) <= 15 else ""


def admin_only(f):
    @wraps(f)
    def w(*a, **k):
        if not session.get("admin"):
            return jsonify(error="Please log in."), 401
        return f(*a, **k)
    return w


# ---------- public ----------
@app.get("/")
def home():
    return render_template("index.html")


@app.get("/api/public")
def public():
    cs = [dict(r) for r in db().execute(
        "select id,name,sector,city,region,state,about,website,phone from companies where active=1 order by name")]
    for c in cs:
        c["wa"] = bool(norm(c.pop("phone")))
    ps = [dict(r) for r in db().execute(
        "select p.id,p.company_id,p.kind,p.title,p.details,p.extra,p.created,c.name company,c.region,c.city,c.sector,c.phone "
        "from posts p left join companies c on c.id=p.company_id and c.active=1 "
        "where p.active=1 and (p.company_id is null or c.id is not null) order by p.id desc")]
    for p in ps:
        p["wa"] = bool(norm(p.pop("phone")))
    return jsonify(companies=cs, posts=ps, community=COMMUNITY_LINK)


@app.get("/go/<int:cid>")
def go(cid):
    r = db().execute("select phone from companies where id=? and active=1", (cid,)).fetchone()
    n = norm(r["phone"]) if r else ""
    if not n:
        abort(404)
    text = request.args.get("t", "Hi, I found you on the Gujarat Accelerate network.")[:300]
    return redirect("https://wa.me/" + n + "?text=" + quote(text))


# ---------- admin ----------
@app.get("/admin")
def admin_page():
    return render_template("admin.html")


@app.post("/api/login")
def login():
    pw = (request.get_json(silent=True) or {}).get("password", "")
    if hmac.compare_digest(pw.encode(), ADMIN_PASSWORD.encode()):
        session["admin"] = True
        return jsonify(ok=True)
    time.sleep(1)
    return jsonify(error="Wrong password."), 401


@app.post("/api/logout")
def logout():
    session.clear()
    return jsonify(ok=True)


@app.get("/api/admin/state")
@admin_only
def state():
    d = db()
    cs = [dict(r) for r in d.execute("select * from companies order by name")]
    ps = [dict(r) for r in d.execute(
        "select p.*,c.name company from posts p left join companies c on c.id=p.company_id where p.active=1 order by p.id desc limit 100")]
    log = [dict(r) for r in d.execute(
        "select p.title,n.post_id,n.status,count(*) n from notif n join posts p on p.id=n.post_id "
        "group by n.post_id,n.status order by n.post_id desc limit 40")]
    return jsonify(companies=cs, posts=ps, log=log, whatsapp_live=bool(WA_TOKEN and WA_PHONE_ID))


def cell(row, h, key):
    i = h.get(key)
    return "" if i is None or row[i] is None else str(row[i]).strip()


def items(s):
    return [x.strip() for x in re.split(r"[;\n]", s) if x.strip()]


@app.post("/api/admin/import")
@admin_only
def import_xlsx():
    f = request.files.get("file")
    if not f:
        return jsonify(error="Choose an Excel file first."), 400
    try:
        wb = load_workbook(f, read_only=True, data_only=True)
        ws = wb["Companies"] if "Companies" in wb.sheetnames else wb.worksheets[0]
        rows = list(ws.iter_rows(values_only=True))
        h = {str(x).strip().lower(): i for i, x in enumerate(rows[0]) if x}
    except Exception:
        return jsonify(error="Could not read that file. Use an .xlsx with a header row."), 400
    if "name" not in h:
        return jsonify(error="The sheet needs a 'name' column in the header row."), 400
    opt_i = next((i for k, i in h.items() if "opt" in k), None)
    d, added, updated, skipped = db(), 0, 0, 0
    for r in rows[1:]:
        name = cell(r, h, "name")
        if not name or name.lower().startswith("example"):
            skipped += 1
            continue
        region = next((x for x in REGIONS if cell(r, h, "region").lower().startswith(x.lower())), "")
        optin = 1 if opt_i is not None and str(r[opt_i] or "").strip().lower() in ("y", "yes", "1", "true") else 0
        vals = (cell(r, h, "sector"), cell(r, h, "city"), region, cell(r, h, "state") or "Gujarat", cell(r, h, "about"),
                norm(cell(r, h, "whatsapp")), cell(r, h, "website"), optin)
        ex = d.execute("select id from companies where name=?", (name,)).fetchone()
        if ex:
            d.execute("update companies set sector=?,city=?,region=?,state=?,about=?,phone=?,website=?,optin=? where id=?", vals + (ex["id"],))
            updated += 1
        else:
            cid = d.execute("insert into companies(sector,city,region,state,about,phone,website,optin,name) values(?,?,?,?,?,?,?,?,?)",
                            vals + (name,)).lastrowid
            for kind, col in (("sell", "sells"), ("buy", "buys"), ("job", "jobs")):
                for t in items(cell(r, h, col)):
                    d.execute("insert into posts(company_id,kind,title,created) values(?,?,?,?)", (cid, kind, t[:150], now()))
            added += 1
    d.commit()
    return jsonify(ok=True, added=added, updated=updated, skipped=skipped)


@app.post("/api/admin/company/<int:cid>/toggle")
@admin_only
def toggle_company(cid):
    d = db()
    d.execute("update companies set active=1-active where id=?", (cid,))
    d.commit()
    return jsonify(ok=True)


@app.post("/api/admin/reset")
@admin_only
def reset_all():
    if (request.get_json(silent=True) or {}).get("confirm") != "DELETE":
        return jsonify(error="Confirmation missing."), 400
    d = db()
    d.executescript("delete from notif; delete from posts; delete from companies;")
    d.commit()
    return jsonify(ok=True)


@app.post("/api/admin/post")
@admin_only
def add_post():
    j = request.get_json(silent=True) or {}
    kind, title = j.get("kind"), (j.get("title") or "").strip()
    if kind not in KINDS or not title:
        return jsonify(error="Choose a type and enter a title."), 400
    d = db()
    pid = d.execute("insert into posts(company_id,kind,title,details,extra,created) values(?,?,?,?,?,?)",
                    (j.get("company_id") or None, kind, title[:150], (j.get("details") or "")[:600], (j.get("extra") or "")[:120], now())).lastrowid
    d.commit()
    if j.get("notify"):
        threading.Thread(target=notify, args=(pid,), daemon=True).start()
    return jsonify(ok=True, id=pid)


@app.post("/api/admin/post/<int:pid>/delete")
@admin_only
def del_post(pid):
    d = db()
    d.execute("update posts set active=0 where id=?", (pid,))
    d.commit()
    return jsonify(ok=True)


@app.post("/api/admin/test")
@admin_only
def test_msg():
    n = norm((request.get_json(silent=True) or {}).get("phone"))
    if not n:
        return jsonify(error="Enter a valid WhatsApp number."), 400
    st, detail = send_wa(n, ["Test message", "If you can read this, the connection works", SITE_URL])
    return jsonify(status=st, detail=detail)


# ---------- WhatsApp ----------
def clean(s):
    return re.sub(r"\s+", " ", str(s or "")).strip()[:300] or "-"


def send_wa(phone, params):
    if not (WA_TOKEN and WA_PHONE_ID):
        return "dry-run", "WhatsApp credentials not set. Nothing was sent."
    body = {"messaging_product": "whatsapp", "to": phone, "type": "template",
            "template": {"name": WA_TEMPLATE, "language": {"code": WA_LANG},
                         "components": [{"type": "body", "parameters": [{"type": "text", "text": clean(p)} for p in params]}]}}
    try:
        r = requests.post(f"https://graph.facebook.com/v21.0/{WA_PHONE_ID}/messages",
                          headers={"Authorization": "Bearer " + WA_TOKEN}, json=body, timeout=20)
        return ("sent" if r.ok else "failed"), r.text[:300]
    except requests.RequestException as e:
        return "failed", str(e)[:300]


def notify(pid):
    c = sqlite3.connect(DB)
    c.row_factory = sqlite3.Row
    p = c.execute("select p.*,c.name company from posts p left join companies c on c.id=p.company_id where p.id=?", (pid,)).fetchone()
    if not p:
        return
    params = [f"{KINDS[p['kind']]}: {p['title']}", " - ".join(x for x in (p["company"], p["extra"]) if x) or "Gujarat Accelerate network",
              f"{SITE_URL}/#{TAB[p['kind']]}"]
    rec = c.execute("select phone from companies where active=1 and optin=1 and phone!='' and id is not ?", (p["company_id"],)).fetchall()
    for r in rec:
        st, detail = send_wa(r["phone"], params)
        c.execute("insert into notif(post_id,phone,status,detail,at) values(?,?,?,?,?)", (pid, r["phone"], st, detail, now()))
        c.commit()
        time.sleep(0.15)
    c.close()


init_db()
if __name__ == "__main__":
    app.run(debug=True, port=5000)
